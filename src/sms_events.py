"""Parse Dialpad SMS webhook events as Pabbly logs them to the Google Sheet.

Each sheet row is one webhook delivery: column `received_at_et` and column `raw_json`.
`raw_json` is usually a JSON string literal that contains the payload JSON (double-encoded),
and message text can hold raw newlines, so decoding is lenient.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from src import phones

EASTERN = ZoneInfo("America/New_York")
MMS_URL_PREFIX = "https://content.dialpad.com/s/img/"
RAW_JSON_HEADER = "raw_json"
RECEIVED_AT_HEADER = "received_at_et"


class SmsParseError(ValueError):
    pass


@dataclass(frozen=True)
class Party:
    phone: str
    name: str


@dataclass(frozen=True)
class SheetRow:
    source: str  # "<tab>!<row number>" for traceability
    received_at: Any
    raw_json: Any


@dataclass(frozen=True)
class SmsEvent:
    event_id: int
    created_at: datetime  # UTC
    direction: str  # "inbound" | "outbound"
    from_number: str | None
    to_numbers: tuple[str, ...]
    text: str
    is_mms: bool
    contact_id: str
    contact_type: str  # "local" | "google" | "shortcode" | "contact_group" | ...
    contact_name: str
    contact_phone: str | None
    members: tuple[Party, ...]  # group threads only; Dialpad omits the target user
    user_phones: tuple[str, ...]  # members that are Dialpad users (member type "user")
    target: Party | None  # the Dialpad line this copy of the event was delivered for
    target_type: str | None
    status: str | None
    is_slim: bool  # trimmed payload without target/members/status
    source: str
    raw: dict[str, Any]

    @property
    def is_group(self) -> bool:
        return self.contact_type == "contact_group"

    @property
    def created_at_et(self) -> datetime:
        return self.created_at.astimezone(EASTERN)


def decode_raw_json(raw: Any) -> dict[str, Any]:
    value = raw
    for _ in range(3):
        if isinstance(value, dict):
            return value
        if not isinstance(value, str):
            break
        try:
            value = json.loads(value, strict=False)
        except json.JSONDecodeError as exc:
            raise SmsParseError(f"raw_json is not valid JSON: {exc}") from exc
    if isinstance(value, dict):
        return value
    raise SmsParseError(f"raw_json did not decode to an object (got {type(value).__name__})")


def parse_event(payload: dict[str, Any], *, source: str = "") -> SmsEvent:
    try:
        event_id = int(payload["id"])
        created_ms = int(payload["created_date"])
        contact = payload["contact"] or {}
    except (KeyError, TypeError, ValueError) as exc:
        raise SmsParseError(f"{source}: missing id/created_date/contact") from exc

    text = payload.get("text") or ""
    mms_url = payload.get("mms_url")
    is_mms = bool(payload.get("mms")) or bool(mms_url) or text.startswith(MMS_URL_PREFIX)

    target_raw = payload.get("target") or None
    target = None
    if target_raw and target_raw.get("phone"):
        target = Party(phones.normalize(target_raw["phone"]), target_raw.get("name") or "")

    members = tuple(
        Party(phones.normalize(m["phone"]), m.get("name") or "")
        for m in contact.get("members") or []
        if m.get("phone")
    )

    user_phones = tuple(
        phones.normalize(m["phone"])
        for m in contact.get("members") or []
        if m.get("phone") and m.get("type") == "user"
    )

    return SmsEvent(
        event_id=event_id,
        created_at=datetime.fromtimestamp(created_ms / 1000, tz=timezone.utc),
        direction=payload.get("direction") or "",
        from_number=phones.normalize(payload.get("from_number")),
        to_numbers=tuple(
            n for n in (phones.normalize(x) for x in payload.get("to_number") or []) if n
        ),
        text=text,
        is_mms=is_mms,
        contact_id=str(contact.get("id") or ""),
        contact_type=contact.get("type") or "",
        contact_name=contact.get("name") or "",
        contact_phone=phones.normalize(contact.get("phone")),
        members=members,
        user_phones=user_phones,
        target=target,
        target_type=(target_raw or {}).get("type"),
        status=payload.get("message_status"),
        is_slim="target" not in payload,
        source=source,
        raw=payload,
    )


def parse_rows(rows: Iterable[SheetRow]) -> tuple[list[SmsEvent], list[str]]:
    """Parse sheet rows into events. Returns (events, errors); a bad row never stops the run."""
    events: list[SmsEvent] = []
    errors: list[str] = []
    for row in rows:
        if row.raw_json in (None, ""):
            continue
        try:
            events.append(parse_event(decode_raw_json(row.raw_json), source=row.source))
        except SmsParseError as exc:
            errors.append(f"{row.source}: {exc}")
    return events, errors


def read_xlsx_rows(path: Path, sheet_names: Iterable[str] | None = None) -> list[SheetRow]:
    """Read rows from an .xlsx export of the Google Sheet (proof-of-concept input)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    wanted = list(sheet_names) if sheet_names else workbook.sheetnames
    rows: list[SheetRow] = []
    for name in wanted:
        if name not in workbook.sheetnames:
            raise SmsParseError(f"{path.name} has no tab named {name!r} (tabs: {workbook.sheetnames})")
        values = workbook[name].iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(values, ())]
        if RAW_JSON_HEADER not in header:
            continue
        raw_col = header.index(RAW_JSON_HEADER)
        received_col = header.index(RECEIVED_AT_HEADER) if RECEIVED_AT_HEADER in header else None
        for offset, row in enumerate(values, start=2):
            raw = row[raw_col] if raw_col < len(row) else None
            received = row[received_col] if received_col is not None and received_col < len(row) else None
            rows.append(SheetRow(source=f"{name}!{offset}", received_at=received, raw_json=raw))
    workbook.close()
    return rows
